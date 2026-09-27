"""
family_results.json -> har bir oila a'zosi uy egasiga kim bo'lishini TAXMIN qiladi
va family_relations.xlsx ga yozadi.

Sayt oila a'zosining holatini (familyMemberStatus) bermaydi, shuning uchun u
quyidagilar bo'yicha aniqlanadi:
  * jinsi       — JSHSHIR ning 1-raqami (toq = erkak, juft = ayol);
  * yosh farqi  — uy egasining tug'ilgan sanasi bilan solishtiriladi;
  * otasining ismi — farzandning otasining ismi = otaning ismi,
                  aka-uka/opa-singilning otasining ismi bir xil;
  * familiyasi  — ko'pincha bobosining ismidan yasaladi (AKRAMOV <- AKRAM).

Har bir taxmin uchun "Ishonch" (Yuqori / O'rta / Past) va "Sabab" yoziladi —
Past va O'rta bo'lganlarini qo'lda tekshirib chiqing.

Ishlatish:
    python guess_relations.py
"""

import datetime
import json
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from add_birthdate import decode_birthdate
from fetch_family import RESULTS_FILE, read_people

OUT_FILE = "family_relations.xlsx"

HIGH, MID, LOW = "Yuqori", "O'rta", "Past"

CHILD_GAP = 14      # ota-ona va farzand orasidagi eng kam yosh farqi
GRAND_GAP = 35      # bobo/buvi va nevara orasidagi eng kam yosh farqi
GRAND_UNSURE = 42   # 35..42 yil farq — farzand ham, nevara ham bo'lishi mumkin
SIBLING_MAX = 25    # aka-uka/opa-singil orasidagi eng katta yosh farqi
SPOUSE_MAX = 20     # er-xotin orasidagi eng katta yosh farqi
PAIR_MAX = 12       # kelin/kuyov/yanga — juftining yoshiga yaqinligi

# --- Ismlarni solishtirish ------------------------------------------------

CYR = {
    "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Ё": "YO", "Ж": "J",
    "З": "Z", "И": "I", "Й": "Y", "К": "K", "Л": "L", "М": "M", "Н": "N",
    "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U", "Ф": "F",
    "Х": "X", "Ц": "TS", "Ч": "CH", "Ш": "SH", "Щ": "SH", "Ъ": "'", "Ы": "I",
    "Ь": "", "Э": "E", "Ю": "YU", "Я": "YA", "Ў": "O'", "Қ": "Q", "Ғ": "G'",
    "Ҳ": "H",
}
VOWELS = set("АЕЁИОУЭЮЯЎЪЬ")


def translit(text):
    """Kirill -> lotin (Е so'z boshida va unlidan keyin YE bo'ladi)."""
    out, prev = [], ""
    for ch in text.upper():
        if ch == "Е":
            out.append("YE" if (not prev or not prev.isalpha() or prev in VOWELS) else "E")
        else:
            out.append(CYR.get(ch, ch))
        prev = ch
    return "".join(out)


def skeleton(word):
    """Solishtirish uchun: X=H, Q=K, apostroflarsiz, qo'sh harflarsiz."""
    s = translit(word)
    s = re.sub(r"[^A-Z]", "", s)
    s = s.replace("X", "H").replace("Q", "K").replace("DJ", "J").replace("YE", "E")
    return re.sub(r"(.)\1+", r"\1", s)


def lev(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def same(a, b):
    """Ikki ism bir xilmi (imlo farqlariga chidamli: JAMOL~JALOL, FERUZ~FERUZJON)."""
    if not a or not b:
        return False
    if a == b:
        return True
    strip = lambda x: x[:-3] if x.endswith("JON") and len(x) > 5 else x
    if strip(a) == strip(b):
        return True
    # VALERIY -> VALERYEVNA (ildizi VALER)
    tail = lambda x: re.sub(r"(IY|I|Y)$", "", x) if len(x) > 4 else x
    if tail(a) == tail(b):
        return True
    n, d = min(len(a), len(b)), lev(a, b)
    return (n >= 4 and d <= 1) or (n >= 7 and d <= 2)


PATR_WORDS = {"OGLI", "UGLI", "KIZI", "KZI"}
PATR_SUFFIXES = ("OVICH", "EVICH", "ICHNA", "OVNA", "EVNA", "ICH", "NA")


def split_name(words):
    """ISM [ISM2] FAMILIYA OTASI [O'G'LI/QIZI] -> (ism, familiya, otasining ismi so'zlari).

    Oxiridan tahlil qilinadi, chunki ism ikki so'zli bo'lishi mumkin
    (MUHAMMAD UMAR OBIDOV SHOHRUHXON O'G'LI)."""
    if len(words) < 3:
        return (words[0] if words else ""), (words[1] if len(words) > 1 else ""), []
    n = 2 if (len(words) >= 4 and skeleton(words[-1]) in PATR_WORDS) else 1
    return " ".join(words[: -n - 1]), words[-n - 1], words[-n:]


def patronymic_root(words):
    """['ANVAROVICH'] / ['ANVAR', "O'G'LI"] -> 'ANVAR' (skelet)."""
    if not words:
        return ""
    if len(words) >= 2 and skeleton(words[-1]) in PATR_WORDS:
        return skeleton(words[-2])
    s = skeleton("".join(words))
    for suf in PATR_SUFFIXES:
        if s.endswith(suf) and len(s) - len(suf) >= 2:
            return s[: -len(suf)]
    return s


def surname_root(word):
    """SULTONOVA -> SULTON, XUDOYEV -> HUDO."""
    s = skeleton(word)
    if s.endswith(("OVA", "EVA", "INA")):
        s = s[:-1]
    for suf in ("OV", "EV", "IN"):
        if s.endswith(suf) and len(s) - len(suf) >= 2:
            return s[: -len(suf)]
    return s


# --- Shaxs ----------------------------------------------------------------

def parse_date(value, pinfl):
    try:
        return datetime.datetime.strptime(str(value), "%d.%m.%Y").date()
    except (ValueError, TypeError):
        return decode_birthdate(pinfl)


class Person:
    def __init__(self, full_name, pinfl, birth_on, surname_first=False):
        self.full_name = full_name or ""
        self.pinfl = str(pinfl or "")
        self.birth_on = birth_on or ""
        words = self.full_name.split()
        if surname_first and len(words) >= 2:  # family_search.xlsx: FAMILIYA ISM OTASI
            words[0], words[1] = words[1], words[0]
        first, surname, patr_words = split_name(words)
        self.first = skeleton(first.split()[0]) if first else ""
        self.surname = surname_root(surname) if surname else ""
        self.patr = patronymic_root(patr_words)
        self.dob = parse_date(birth_on, self.pinfl)
        if self.pinfl[:1].isdigit() and self.pinfl[0] != "0":
            self.male = int(self.pinfl[0]) % 2 == 1
        else:
            self.male = any(skeleton(w) in {"OGLI", "UGLI"} or skeleton(w).endswith("VICH")
                            for w in patr_words)
        self.diff = 0.0  # uy egasidan necha yosh katta (+) / kichik (-)
        self.rel = self.conf = self.why = None

    def set(self, rel, conf, why):
        if self.rel is None:
            self.rel, self.conf, self.why = rel, conf, why
            return True
        return False

    def age(self, today):
        if not self.dob:
            return None
        return today.year - self.dob.year - ((today.month, today.day) < (self.dob.month, self.dob.day))


def g(p, male_word, female_word):
    return male_word if p.male else female_word


def sibling_word(p):
    return g(p, "Akasi" if p.diff > 0 else "Ukasi", "Opasi" if p.diff > 0 else "Singlisi")


# --- Asosiy mantiq --------------------------------------------------------

def infer(owner, members):
    """members ichida owner ham bor. Har biriga rel/conf/why yoziladi."""
    rest = [m for m in members if m is not owner]
    for m in members:
        if owner.dob and m.dob:
            m.diff = (owner.dob - m.dob).days / 365.25
    owner.set("O'zi (qidirilgan shaxs)", HIGH, "JSHSHIR bir xil")

    free = lambda: [m for m in rest if m.rel is None]
    by_rel = lambda *rels: [m for m in rest if m.rel in rels]
    older = lambda m, other, gap: m.diff - other.diff >= gap  # m otherdan gap yil katta

    # 1. Aka-uka / opa-singil: otasining ismi bir xil
    for m in free():
        if owner.patr and same(m.patr, owner.patr) and abs(m.diff) < SIBLING_MAX:
            m.set(sibling_word(m), HIGH, "otasining ismi uy egasiniki bilan bir xil")

    # 2. Farzand: otasining ismi = uy egasining ismi (uy egasi erkak)
    for m in free():
        if owner.male and m.diff <= -CHILD_GAP and same(m.patr, owner.first):
            m.set(g(m, "O'g'li", "Qizi"), HIGH, "otasining ismi = uy egasining ismi")

    # 3. Ota: ismi = uy egasining otasining ismi
    for m in free():
        if m.male and m.diff >= CHILD_GAP and same(m.first, owner.patr):
            m.set("Otasi", HIGH, "ismi = uy egasining otasining ismi")
    if not any(m.rel == "Otasi" for m in rest):
        for m in free():
            if m.male and m.diff >= CHILD_GAP and same(m.patr, owner.surname):
                m.set("Otasi", MID, "otasining ismi uy egasining familiyasiga asos bo'lgan")
                break

    # 4. Turmush o'rtog'i: qarama-qarshi jins, yoshi yaqin
    cands = [m for m in free() if m.male != owner.male and abs(m.diff) <= SPOUSE_MAX]
    spouse = None
    if cands:
        def kids_of(man):
            return [k for k in rest if k.diff <= man.diff - CHILD_GAP and same(k.patr, man.first)]

        def score(m):
            father = m if m.male else owner
            return len(kids_of(father)) * 10 - abs(m.diff)

        spouse = max(cands, key=score)
        father = spouse if spouse.male else owner
        kids = kids_of(father)
        word = g(spouse, "Eri", "Xotini")
        if spouse.male and kids:
            spouse.set(word, HIGH, "farzandlarning otasining ismi = uning ismi")
        elif kids and len(cands) == 1 and all(older(spouse, k, CHILD_GAP) for k in kids):
            spouse.set(word, HIGH, "yoshi yaqin, farzandlarga ona bo'la oladi")
        else:
            spouse.set(word, MID, "qarama-qarshi jins, yoshi yaqin")

    # 5. Turmush o'rtog'i orqali: farzandlar, qaynota, qaynisi/qaynsinglisi
    if spouse:
        for m in free():
            if spouse.male and older(spouse, m, CHILD_GAP) and same(m.patr, spouse.first):
                m.set(g(m, "O'g'li", "Qizi"), HIGH, "otasining ismi = erining ismi")
        for m in free():
            if m.male and older(m, spouse, CHILD_GAP) and same(m.first, spouse.patr):
                m.set("Qaynotasi", HIGH, "ismi = turmush o'rtog'ining otasining ismi")
        for m in free():
            if spouse.patr and same(m.patr, spouse.patr) and abs(m.diff - spouse.diff) < SIBLING_MAX:
                m.set(g(m, "Qaynisi (turmush o'rtog'ining aka-ukasi)",
                        "Qaynsinglisi (turmush o'rtog'ining opa-singlisi)"),
                      HIGH, "otasining ismi turmush o'rtog'iniki bilan bir xil")

    def link_children():
        """Nevaralar va jiyanlar: otasi — o'g'il / kuyov / aka-uka / pochcha."""
        changed = True
        while changed:
            changed = False
            for m in free():
                for parent_rels, rel in (
                    (("O'g'li", "Kuyovi"), "Nevarasi"),
                    (("Akasi", "Ukasi", "Pochchasi (opa/singlisining eri)"), "Jiyani"),
                ):
                    for par in by_rel(*parent_rels):
                        if par.male and older(par, m, CHILD_GAP) and same(m.patr, par.first):
                            why = (f"otasi ({par.full_name.split()[0]}) — "
                                   f"uy egasining {par.rel.split()[0].lower()}")
                            changed |= m.set(rel, HIGH, why)
                            break
                    if m.rel:
                        break

    link_children()

    # 6. Bobo: ismi = otasining otasining ismi yoki uy egasining familiyasi
    father = next(iter(by_rel("Otasi")), None)
    for m in free():
        if m.male and m.diff >= GRAND_GAP:
            if father and same(m.first, father.patr):
                m.set("Bobosi", HIGH, "ismi = otasining otasining ismi")
            elif same(m.first, owner.surname):
                m.set("Bobosi", MID, "ismi uy egasining familiyasiga asos bo'lgan")

    # 7. Ona / qaynona: otasi / qaynotasi bilan tengdosh ayol
    father_in_law = next(iter(by_rel("Qaynotasi")), None)
    for m in free():
        if not m.male and m.diff >= CHILD_GAP:
            if father and abs(m.diff - father.diff) <= SPOUSE_MAX:
                m.set("Onasi", HIGH, "otasi bilan tengdosh")
            elif father_in_law and abs(m.diff - father_in_law.diff) <= SPOUSE_MAX:
                m.set("Qaynonasi", HIGH, "qaynotasi bilan tengdosh")

    # 8. Kelin / kuyov / yanga / pochcha: o'g'il/qiz/aka/opaning juftiga yaqin yosh
    for m in free():
        pairs = (
            (not m.male, ("O'g'li",), "Kelini", "o'g'lining yoshiga yaqin"),
            (m.male, ("Qizi",), "Kuyovi", "qizining yoshiga yaqin"),
            (not m.male, ("Akasi", "Ukasi"), "Yangasi (aka/ukasining xotini)",
             "aka/ukasining yoshiga yaqin"),
            (m.male, ("Opasi", "Singlisi"), "Pochchasi (opa/singlisining eri)",
             "opa/singlisining yoshiga yaqin"),
        )
        for ok, rels, rel, why in pairs:
            if ok and any(abs(m.diff - p.diff) <= PAIR_MAX and not same(m.patr, p.patr)
                          for p in by_rel(*rels)):
                m.set(rel, MID, why)
                break

    link_children()

    # 9. Qolganlar — faqat yosh va jins bo'yicha
    for m in free():
        if m.diff >= GRAND_GAP:
            m.set(g(m, "Bobosi", "Buvisi"), LOW, f"uy egasidan {m.diff:.0f} yosh katta")
        elif m.diff >= CHILD_GAP:
            if m.male:
                conf = MID if same(m.surname, owner.surname) else LOW
                m.set("Otasi", conf, f"uy egasidan {m.diff:.0f} yosh katta erkak")
            elif (not owner.male and spouse and same(m.surname, spouse.surname)
                  and not same(m.surname, owner.surname)):
                m.set("Qaynonasi", MID, "familiyasi turmush o'rtog'iniki bilan bir xil")
            else:
                m.set("Onasi", MID, f"uy egasidan {m.diff:.0f} yosh katta ayol")
        elif m.diff <= -GRAND_UNSURE:
            m.set("Nevarasi", MID, f"uy egasidan {-m.diff:.0f} yosh kichik")
        elif m.diff <= -GRAND_GAP:
            rel = g(m, "O'g'li", "Qizi") if same(m.surname, owner.surname) else "Nevarasi"
            m.set(rel, LOW, f"{-m.diff:.0f} yosh kichik: farzand ham, nevara ham bo'lishi mumkin")
        elif m.diff <= -CHILD_GAP:
            if not owner.male:
                m.set(g(m, "O'g'li", "Qizi"), MID,
                      "uy egasi — onasi (otasi oila tarkibida yo'q)")
            elif spouse and older(spouse, m, CHILD_GAP):
                m.set(g(m, "O'gay o'g'li", "O'gay qizi"), LOW,
                      "otasining ismi uy egasiniki emas, xotinining farzandi bo'lishi mumkin")
            else:
                m.set("Qarindoshi (farzand yoshida)", LOW,
                      "otasining ismi uy egasiniki emas")
        elif same(m.surname, owner.surname):
            m.set(sibling_word(m) + " (taxminiy)", LOW, "familiyasi bir xil, yoshi yaqin")
        else:
            m.set("Qarindoshi (aniqlanmadi)", LOW, "yoshi yaqin, bog'liqlik topilmadi")


# --- Excel ----------------------------------------------------------------

def build_rows(people, results, today):
    rows, stats = [], Counter()
    for p in people:
        res = results.get(p["pinfl"])
        head = [p["num"], p["name"], p["pinfl"], p["dob"]]
        if not res or res["status"] != "ok" or not res["members"]:
            note = ("Hali so'ralmagan" if not res else
                    res.get("error") if res["status"] != "ok" else "Oila a'zolari topilmadi")
            rows.append((head + [""] * 8 + [note], None))
            continue

        members = [Person(m["fullName"], m["pinfl"], m["birthOn"]) for m in res["members"]]
        owner = next((m for m in members if m.pinfl == p["pinfl"]), None)
        if owner is None:  # uy egasi ro'yxatda bo'lmasa — Excel'dagi ismidan
            owner = Person(p["name"], p["pinfl"], p["dob"], surname_first=True)
            members.insert(0, owner)
        infer(owner, members)

        others = sorted((m for m in members if m is not owner), key=lambda m: -m.diff)
        for m in [owner] + others:
            stats[(m.rel, m.conf)] += 1
            rows.append((head + [
                m.full_name, m.pinfl, m.birth_on or (m.dob.strftime("%d.%m.%Y") if m.dob else ""),
                "Erkak" if m.male else "Ayol", m.age(today),
                "" if m is owner else f"{m.diff:+.0f}", m.rel, m.conf, m.why,
            ], m.conf))
    return rows, stats


def write_excel(rows, stats):
    wb = Workbook()
    ws = wb.active
    ws.title = "Oila a'zolari"
    headers = ["№", "Uy egasi (F.I.O.)", "Uy egasi JSHSHIR", "Uy egasi tug'ilgan sanasi",
               "Oila a'zosi (F.I.O.)", "JSHSHIR", "Tug'ilgan sana", "Jinsi", "Yoshi",
               "Yosh farqi", "Taxminiy holat", "Ishonch", "Sabab"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    group_fill = PatternFill("solid", fgColor="EEF3F8")
    conf_style = {
        HIGH: (PatternFill("solid", fgColor="D5F0DD"), Font(bold=True, color="1E7B3A")),
        MID: (PatternFill("solid", fgColor="FFF2CC"), Font(bold=True, color="8A6D00")),
        LOW: (PatternFill("solid", fgColor="FDE9E7"), Font(bold=True, color="C0392B")),
    }
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border

    group, prev_owner = -1, None
    for line, conf in rows:
        if (line[0], line[2]) != prev_owner:
            group, prev_owner = group + 1, (line[0], line[2])
        ws.append(line)
        r = ws.max_row
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=r, column=col)
            c.border = border
            c.alignment = left if col in (2, 5, 11, 13) else center
            if group % 2:
                c.fill = group_fill
        for col in (3, 4, 6, 7):
            ws.cell(row=r, column=col).number_format = "@"
        if conf:
            fill, font = conf_style[conf]
            ws.cell(row=r, column=12).fill = fill
            ws.cell(row=r, column=12).font = font
            ws.cell(row=r, column=11).font = Font(bold=True)
        elif line[-1]:
            ws.cell(row=r, column=13).font = Font(color="C0392B")

    for idx, w in enumerate([6, 34, 17, 13, 34, 17, 13, 8, 7, 8, 30, 10, 48], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 34
    ws.sheet_view.showGridLines = False

    # --- Xulosa varag'i ---
    ss = wb.create_sheet("Xulosa")
    ss.append(["Taxminiy holat", HIGH, MID, LOW, "Jami"])
    by_rel = Counter()
    for (rel, _), n in stats.items():
        by_rel[rel] += n
    for rel, total in by_rel.most_common():
        ss.append([rel] + [stats.get((rel, c), 0) for c in (HIGH, MID, LOW)] + [total])
    ss.append(["Jami"] + [sum(n for (_, c), n in stats.items() if c == k) for k in (HIGH, MID, LOW)]
              + [sum(stats.values())])
    for cell in ss[1]:
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
    for row in ss.iter_rows(min_row=2):
        for c in row:
            c.border = border
            c.alignment = left if c.column == 1 else center
    for c in ss[ss.max_row]:
        c.font = Font(bold=True)
    ss.column_dimensions["A"].width = 44
    for col in "BCDE":
        ss.column_dimensions[col].width = 11

    wb.save(OUT_FILE)


def main():
    people = read_people()
    with open(RESULTS_FILE, encoding="utf-8") as f:
        results = json.load(f)
    rows, stats = build_rows(people, results, datetime.date.today())
    write_excel(rows, stats)

    by_conf = Counter()
    for (_, c), n in stats.items():
        by_conf[c] += n
    print(f"Tayyor: {OUT_FILE}  ({len(rows)} qator)")
    print(f"Ishonch: Yuqori {by_conf[HIGH]}, O'rta {by_conf[MID]}, Past {by_conf[LOW]}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
