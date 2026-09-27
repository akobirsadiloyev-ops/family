"""
yakuniy_natija.xlsx bo'yicha qisqa hisobot. Ekranga FAQAT SONLAR chiqadi —
ism, JSHSHIR, telefon chiqmaydi, shuning uchun natijani bemalol ulashsa bo'ladi.

Ko'rsatadi:
  * xonadonlar holati ("Oila holati (670)" varag'i)
  * oila a'zolari: qarindoshlik turi, ishonch darajasi, manba bo'yicha soni
    ("Oila a'zolari (taxminiy)" varag'i)

Qarindoshligi bo'sh yoki ishonchi O'rta/Past bo'lgan a'zolar qo'lda tekshirish
uchun tekshirish_kerak.xlsx ga yoziladi (u kompyuterda qoladi, GitHub'ga ketmaydi).

Ishlatish:
    python hisobot.py
"""

import sys
from collections import Counter

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# --- Sozlamalar ----------------------------------------------------------
FILE = "yakuniy_natija.xlsx"
HOLAT_SHEET = "Oila holati (670)"
AZO_SHEET = "Oila a'zolari (taxminiy)"
OUT = "tekshirish_kerak.xlsx"
TEKSHIR_ISHONCH = ("O'rta", "Past")
# -------------------------------------------------------------------------


def read_rows(wb, name):
    """Varaqni [{sarlavha: qiymat}, ...] ko'rinishida qaytaradi."""
    if name not in wb.sheetnames:
        sys.exit(f"'{name}' varag'i topilmadi. Avval: python build_oila.py")
    rows = wb[name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    return header, [dict(zip(header, r)) for r in rows if any(v is not None for v in r)]


def show(title, counter, total=None):
    print(f"\n{title}")
    width = max((len(str(k)) for k in counter), default=0)
    for key, cnt in counter.most_common():
        pct = f"  ({cnt * 100 / total:.0f}%)" if total else ""
        print(f"  {str(key):<{width}}  {cnt:>5}{pct}")


def main():
    try:
        wb = openpyxl.load_workbook(FILE, data_only=True, read_only=True)
    except FileNotFoundError:
        sys.exit(f"{FILE} topilmadi. Skriptni loyiha papkasida ishga tushiring.")
    except PermissionError:
        sys.exit(f"{FILE} ochiq turibdi. Excel'da yopib, qayta urinib ko'ring.")

    # ===== 1) Xonadonlar =====
    _, homes = read_rows(wb, HOLAT_SHEET)
    print(f"=== Xonadonlar: {len(homes)} ta ===")
    show("Holat:", Counter(h.get("Holat") or "(bo'sh)" for h in homes))

    # ===== 2) Oila a'zolari =====
    # Uy egasi ma'lumoti faqat har uyning birinchi qatorida bor — pastga to'ldiramiz.
    header, members = read_rows(wb, AZO_SHEET)
    owner_cols = ("№", "Uy egasi (F.I.O.)", "Uy egasi JSHSHIR", "Uy ID")
    last = {}
    for m in members:
        if m.get("Uy ID") is not None:
            last = {c: m.get(c) for c in owner_cols}
        m.update({c: m.get(c) if m.get(c) is not None else last.get(c) for c in owner_cols})
    wb.close()

    total = len(members)
    uylar = {m["Uy ID"] for m in members}
    print(f"\n=== Oila a'zolari: {len(uylar)} ta uyda, jami {total} ta a'zo ===")
    show("Qarindoshlik turi:",
         Counter(m.get("Taxminiy qarindoshlik") or "(aniqlanmagan)" for m in members), total)
    show("Ishonch darajasi:",
         Counter(m.get("Ishonch") or "(ko'rsatilmagan)" for m in members), total)
    show("Manba:", Counter(m.get("Manba") or "(bo'sh)" for m in members), total)

    # ===== 3) Qo'lda tekshirish kerak bo'lganlar =====
    check = [m for m in members
             if not m.get("Taxminiy qarindoshlik") or m.get("Ishonch") in TEKSHIR_ISHONCH]
    print(f"\nQo'lda tekshirish kerak: {len(check)} ta a'zo, "
          f"{len({m['Uy ID'] for m in check})} ta uyda")
    if not check:
        return 0

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Tekshirish kerak"
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for m in check:
        ws.append([m.get(h) for h in header])
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        out.save(OUT)
    except PermissionError:
        sys.exit(f"{OUT} ochiq turibdi. Excel'da yopib, qayta urinib ko'ring.")
    print(f"Ro'yxat: {OUT}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
