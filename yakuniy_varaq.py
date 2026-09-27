"""
yakuniy_natija.xlsx ga 2 ta varaq qo'shadi (bor bo'lsa, qayta yasaydi):
  1) "Yakuniy xulosa" — jami sonlar: uy egalari, oila a'zolari, qarindoshlik turlari.
  2) "Yakuniy"        — har bir xonadon: uy egasi, oila a'zolari va izoh
                        (uy egasi har uyda BIR MARTA):
       * oila topilgan         -> a'zolar ro'yxati, "Oila topilgan: N ta a'zo"
       * egasi ikkinchi uyda   -> "Oila a'zolari Uy ID ... da ko'rsatilgan"
       * yolg'iz               -> "Qarindoshlari yo'q (bazada faqat o'zi chiqdi)"
       * bazada yo'q           -> "Bazada yo'q (qidiruv natija bermadi)"
     Ishonchi O'rta/Past a'zolar yoniga "Qo'lda tekshirish kerak" yoziladi.

Manba: shu fayldagi "Oila holati (670)" va "Oila a'zolari (taxminiy)" varaqlari —
build_yakuniy.py yoki build_oila.py qayta ishga tushsa, bu skriptni ham qayta
ishga tushiring. Saqlashdan oldin fayl nusxasi yakuniy_natija_zaxira.xlsx ga olinadi.

Ishlatish:
    python yakuniy_varaq.py
"""

import shutil
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from hisobot import AZO_SHEET, FILE, HOLAT_SHEET, TEKSHIR_ISHONCH, read_rows

# --- Sozlamalar ----------------------------------------------------------
BACKUP = "yakuniy_natija_zaxira.xlsx"
SHEET = "Yakuniy"
XULOSA = "Yakuniy xulosa"
FAM, YOLGIZ, YOQ = "Oila topilgan", "Yolg'iz (oila a'zosi yo'q)", "Tarkibi topilmadi (bazada yo'q)"
IZOH = {
    YOLGIZ: "Qarindoshlari yo'q (bazada faqat o'zi chiqdi)",
    YOQ: "Bazada yo'q (qidiruv natija bermadi)",
}
TEKSHIR = "Qo'lda tekshirish kerak"
# -------------------------------------------------------------------------

HF = PatternFill("solid", fgColor="1F4E78"); HFONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9"); GRP = Side(style="medium", color="9CB3D0")
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
LFT = Alignment(horizontal="left", vertical="center")
COLORS = {  # izoh boshlanishi -> (fon, matn)
    "Oila topilgan:": ("D5F0DD", "1E7B3A"),
    "Qarindoshlari yo'q": ("FFF2CC", "8A6D00"),
    "Bazada yo'q": ("FDE9E7", "C0392B"),
    TEKSHIR: ("FCE4D6", "C55A11"),
}
GRAY = ("EEEEEE", "555555")

HEADER = ["№", "Uy ID", "Uy egasi (F.I.O.)", "Uy egasi JSHSHIR", "A'zolar soni",
          "Oila a'zosi (F.I.Sh.)", "A'zo JSHSHIR", "Tug'ilgan sana", "Qarindoshligi",
          "Ishonch", "Izoh"]
WIDTHS = [5, 9, 30, 16, 8, 30, 16, 12, 26, 9, 44]
TEXT_COLS = (2, 4, 7)   # JSHSHIR va ID lar matn bo'lib qolsin


def load():
    try:
        src = openpyxl.load_workbook(FILE, data_only=True, read_only=True)
    except FileNotFoundError:
        sys.exit(f"{FILE} topilmadi. Skriptni loyiha papkasida ishga tushiring.")
    _, homes = read_rows(src, HOLAT_SHEET)
    _, rows = read_rows(src, AZO_SHEET)
    src.close()

    # Oilalar: uy egasi JSHSHIR -> {"uy": Uy ID, "azolar": [...]}
    # (egasi ma'lumoti faqat har guruhning birinchi qatorida bor)
    families, owner = {}, None
    for r in rows:
        if r.get("Uy egasi JSHSHIR") is not None:
            owner = str(r["Uy egasi JSHSHIR"]).strip()
            families.setdefault(owner, {"uy": str(r.get("Uy ID")), "azolar": []})
        if owner:
            families[owner]["azolar"].append(r)
    return homes, families


def none_or_str(v):
    return None if v is None else str(v).strip()


def build(homes, families):
    """Yakuniy varaq qatorlari va xulosa uchun sonlar."""
    by_owner = defaultdict(list)
    for h in homes:
        by_owner[str(h.get("JSHSHIR")).strip()].append(h)
    # Egasi bir nechta uyda bo'lsa, oila Uy ID si mos kelgan uyda (bo'lmasa birinchisida) ko'rsatiladi
    target = {}
    for p, fam in families.items():
        if p in by_owner:
            hs = by_owner[p]
            target[p] = next((h for h in hs if str(h.get("Uy ID")) == fam["uy"]), hs[0])

    rank = {FAM: 0, YOLGIZ: 1, YOQ: 2}
    ordered = sorted(homes, key=lambda h: rank.get(h.get("Holat"), 3))  # barqaror: asl tartib saqlanadi

    out, shown, stat = [], [], Counter()
    for i, h in enumerate(ordered, 1):
        p = str(h.get("JSHSHIR")).strip()
        uy = h.get("Uy ID")
        head = [i, None if uy is None else str(uy), h.get("Uy egasi (F.I.O.)"), p]
        fam = families.get(p)
        holat = h.get("Holat")
        stat["uy"] += 1
        if fam and target.get(p) is h:
            azolar = fam["azolar"]
            stat["oila"] += 1
            stat["azo"] += len(azolar)
            shown.extend(azolar)
            for j, m in enumerate(azolar):
                check = not m.get("Taxminiy qarindoshlik") or m.get("Ishonch") in TEKSHIR_ISHONCH
                stat["tekshir"] += check
                izoh = f"Oila topilgan: {len(azolar)} ta a'zo" if j == 0 else ""
                if check:
                    izoh = f"{izoh}; {TEKSHIR}" if izoh else TEKSHIR
                out.append(((head + [len(azolar)]) if j == 0 else [None] * 5) + [
                    m.get("A'zo (F.I.Sh.)"), none_or_str(m.get("A'zo JSHSHIR")), m.get("A'zo sana"),
                    m.get("Taxminiy qarindoshlik"), m.get("Ishonch"), izoh])
            continue
        if fam:
            stat["ikkinchi_uy"] += 1
            izoh = (f"Oila a'zolari Uy ID {target[p].get('Uy ID')} da ko'rsatilgan "
                    f"(uy egasi bir xil)")
        elif holat in IZOH:
            stat[holat] += 1
            izoh = IZOH[holat]
        else:
            stat["boshqa"] += 1
            izoh = (f"{holat}, lekin a'zolar ro'yxatida yo'q" if holat == FAM
                    else holat or "Holat ko'rsatilmagan")
        out.append(head + [0, None, None, None, None, None, izoh])

    tashqi = [p for p in families if p not in by_owner]
    return out, shown, stat, tashqi


def write_yakuniy(wb, rows):
    ws = wb.create_sheet(SHEET, 1)
    ws.append(HEADER)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HF, HFONT, CTR
        c.border = Border(THIN, THIN, THIN, THIN)
    for row in rows:
        ws.append(row)
        rr = ws.max_row
        first = row[0] is not None
        for c in range(1, len(HEADER) + 1):
            cell = ws.cell(rr, c)
            top = GRP if (first and rr > 2) else THIN
            cell.border = Border(left=THIN, right=THIN, top=top, bottom=THIN)
            cell.alignment = LFT if c in (3, 6, 9, 11) else CTR
            if c in TEXT_COLS:
                cell.number_format = "@"
        izoh = row[-1] or ""
        key = next((k for k in COLORS if izoh.startswith(k)),
                   TEKSHIR if TEKSHIR in izoh else None)
        if izoh:
            fill, color = COLORS[key] if key else GRAY
            ws.cell(rr, len(HEADER)).fill = PatternFill("solid", fgColor=fill)
            ws.cell(rr, len(HEADER)).font = Font(bold=True, color=color)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{ws.max_row}"
    ws.sheet_view.showGridLines = False


def write_xulosa(wb, lines):
    ws = wb.create_sheet(XULOSA, 0)
    ws.append(["Ko'rsatkich", "Soni"])
    for c in ws[1]:
        c.fill, c.font, c.alignment = HF, HFONT, CTR
    for name, value in lines:
        ws.append([name, value])
        if value is None and name:          # bo'lim sarlavhasi
            ws.cell(ws.max_row, 1).font = Font(bold=True, color="1F4E78")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 10
    ws.sheet_view.showGridLines = False


def main():
    homes, families = load()
    rows, shown, stat, tashqi = build(homes, families)

    summary = [
        ("Jami xonadonlar", stat["uy"]),
        ("Oila a'zolari topilgan uy egalari", stat["oila"]),
        ("Jami oila a'zolari (uy egalarisiz)", stat["azo"]),
        ("Uy egalari bilan birga, jami kishi", stat["oila"] + stat["azo"]),
        ("Egasi ikkinchi uyda ko'rsatilgan uylar", stat["ikkinchi_uy"]),
        ("Qarindoshlari yo'q", stat[YOLGIZ]),
        ("Bazada yo'q", stat[YOQ]),
        ("Boshqa holat", stat["boshqa"]),
        ("Qo'lda tekshirish kerak (a'zolar)", stat["tekshir"]),
    ]
    if tashqi:
        summary.append(("Uylar ro'yxatida yo'q egalar (varaqqa kirmadi)", len(tashqi)))
    lines = summary + [
        ("", None),
        ("Qarindoshlik turi", None),
        *Counter(m.get("Taxminiy qarindoshlik") or "(aniqlanmagan)" for m in shown).most_common(),
        ("", None),
        ("Ishonch darajasi", None),
        *Counter(m.get("Ishonch") or "(ko'rsatilmagan)" for m in shown).most_common(),
    ]

    try:
        shutil.copy2(FILE, BACKUP)
        wb = openpyxl.load_workbook(FILE)
        for old in (SHEET, XULOSA):
            if old in wb.sheetnames:
                del wb[old]
        write_xulosa(wb, lines)
        write_yakuniy(wb, rows)
        for ws in wb.worksheets:           # bir nechta varaq birga tanlanib qolmasin
            ws.sheet_view.tabSelected = False
        wb.active = 0
        wb.active.sheet_view.tabSelected = True
        wb.save(FILE)
    except PermissionError:
        sys.exit(f"{FILE} ochiq turibdi. Excel'da yopib, qayta urinib ko'ring.")

    print(f"{FILE} yangilandi: '{XULOSA}' va '{SHEET}' varaqlari (zaxira: {BACKUP})\n")
    width = max(len(n) for n, _ in summary)
    for name, value in summary:
        print(f"  {name:<{width}}  {value:>5}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
