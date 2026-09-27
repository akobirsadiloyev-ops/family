"""
base.xlsx dagi oila a'zolarini online-mahalla.uz ga (survey_homes_family) qo'shadi.

Ustunlar (base.xlsx / "Oila a'zolari"):
    C = Uy egasi JSHSHIR,  E = Uy egasi ID (uy id),
    F = A'zo JSHSHIR,      G = Tug'ilgan sana,  H = Taxminiy holat (qarindoshlik)

Har bir a'zo uchun:
    1) member-can-add/check  -> shu uyda allaqachon bo'lsa (status=true) -> skip (dublikat)
    2) gcp/pinfl             -> F.I.Sh., hujjat seriya/raqami, hujjat turi
    3) saqlash (POST)        -> 201 = muvaffaqiyat; xato bo'lsa skip + izoh (masalan
                               "boshqa xonadonga biriktirilgan")

Xususiyatlari:
    * QAYTA ISHGA TUSHIRSA BO'LADI — add_progress.json da holat saqlanadi, tayyorlar
      o'tkazib yuboriladi (dublikat bo'lmaydi).
    * Natija natija.xlsx ga yoziladi: har qatorda "Natija" ustuni (yashil=qo'shildi,
      ko'k=allaqachon bor, qizil=xato + sabab).

Ishlatish:
    python add_family.py            # SINOV: faqat dastlabki 3 ta (xavfsiz)
    python add_family.py 20         # dastlabki 20 ta yangi a'zo
    python add_family.py all        # HAMMASI

Token eskirsa (401): python get_token.py  -> keyin add_family.py ni qayta ishga tushiring.
"""

import datetime
import json
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.request

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Sozlamalar ----------------------------------------------------------
BASE_IN = "base.xlsx"
SHEET = "Oila a'zolari"
OUT = "natija.xlsx"
PROGRESS = "add_progress.json"
TOKEN_FILE = "token.txt"
DEFAULT_TOKEN = "82726e4e-d2c7-4217-b127-4edea9bf6300"

API = "https://api.online-mahalla.uz"
PARAMS = {"obl_id": 6, "area_id": 608, "district_id": 608035, "street_id": 60800147}
DELAY = 0.3          # a'zolar orasidagi pauza (soniya) — saqlashning o'zi sekin
GCP_RETRY = 3        # tezlik cheklovi/tarmoq xatosida qayta urinish
EXCEL_EVERY = 25     # har shuncha a'zodan keyin natija.xlsx yangilanadi
# -------------------------------------------------------------------------

# Qarindoshlik (base.xlsx H) -> relationship id
REL_MAP = {
    "otasi": 1, "onasi": 2, "opasi": 3, "akasi": 4, "singlisi": 5, "ukasi": 6,
    "eri": 7, "xotini": 7, "turmushortogi": 7, "qizi": 8, "ogli": 9,
    "qaynotasi": 10, "qaynonasi": 11, "kelini": 12, "boshqa": 13, "kuyovi": 14,
    "nevarasi": 15, "nabira": 15,
}

# Ta'lim (study_level_id): 1=Oliy, 2=Tugallanmagan oliy, 3=O'rta maxsus, 4=O'rta, 5=Ma'lumotsiz
# 18+ uchun taqsimot: 50% O'rta maxsus(3), 30% Oliy(1), 20% O'rta(4) — JSHSHIR bo'yicha barqaror
STUDY_PATTERN = ["3", "3", "3", "3", "3", "1", "1", "1", "4", "4"]

TODAY = datetime.date(2026, 9, 26)


def norm_rel(s):
    return re.sub(r"[^a-z]", "", str(s or "").lower())


def load_token():
    if os.path.exists(TOKEN_FILE):
        t = open(TOKEN_FILE, encoding="utf-8").read().strip()
        if t:
            return t
    return DEFAULT_TOKEN


TOKEN = load_token()
HEADERS = {
    "Authorization": f"Bearer {TOKEN}",
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json",
    "Origin": "https://online-mahalla.uz",
    "Referer": "https://online-mahalla.uz/",
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/154.0.0.0 Safari/537.36"),
}


class AuthError(Exception):
    pass


def refresh_token():
    """401 bo'lganda get_token.py orqali tokenni yangilaydi."""
    print(">>> Token yangilanmoqda (get_token.py)...")
    try:
        subprocess.run([sys.executable, "get_token.py"], timeout=180)
    except Exception as e:
        print("   token yangilashda xato:", e)
    global TOKEN
    TOKEN = load_token()
    HEADERS["Authorization"] = f"Bearer {TOKEN}"
    print("   yangi token:", TOKEN[:12], "...")


def _req(method, url, body=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        txt = e.read().decode("utf-8", "replace")
        if e.code in (401, 403):
            raise AuthError(f"HTTP {e.code}")
        return e.code, txt
    except urllib.error.URLError as e:
        return 0, f"URLError: {e.reason}"


def decode_birth(pinfl):
    p = str(pinfl).strip()
    if len(p) != 14 or not p.isdigit():
        return None
    c = int(p[0]); dd, mm, yy = int(p[1:3]), int(p[3:5]), int(p[5:7])
    base = {1: 1800, 2: 1800, 3: 1900, 4: 1900, 5: 2000, 6: 2000}.get(c)
    if base is None:
        return None
    try:
        return datetime.date(base + yy, mm, dd)
    except ValueError:
        return None


def age_of(pinfl):
    d = decode_birth(pinfl)
    if not d:
        return None
    return TODAY.year - d.year - ((TODAY.month, TODAY.day) < (d.month, d.day))


def study_level(age, pinfl):
    if age is None:
        return "4"
    if age <= 7:
        return "5"          # Ma'lumotsiz
    if age <= 17:
        return "4"          # O'rta
    return STUDY_PATTERN[int(pinfl) % 10]


def deep_message(txt):
    """Server javobidan eng ma'noli xabarni chiqaradi."""
    try:
        j = json.loads(txt)
    except Exception:
        return txt[:200]
    best = None

    def walk(o):
        nonlocal best
        if isinstance(o, dict):
            for k in ("message", "error", "result_message", "detail"):
                v = o.get(k)
                if isinstance(v, str) and v.strip() and v.strip().lower() not in ("null", "ok"):
                    best = v.strip()
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
    walk(j)
    return best or txt[:200]


# --- API chaqiruvlari -----------------------------------------------------

_home_cache = {}


def get_home(home_id):
    """Uy id -> (survey_uuid, owner_pinfl, owner_phone). Keshlanadi."""
    if home_id in _home_cache:
        return _home_cache[home_id]
    q = "&".join(f"{k}={v}" for k, v in PARAMS.items())
    url = f"{API}/web/v1/forms/survey_homes/{home_id}?{q}"
    code, txt = _req("GET", url)
    if code != 200:
        raise RuntimeError(f"uy formasi olinmadi (HTTP {code}): {deep_message(txt)}")
    fd = (json.loads(txt).get("formData") or {})
    info = (str(fd.get("survey_uuid") or ""), str(fd.get("pinfl") or ""),
            str(fd.get("mobile_phone") or ""))
    _home_cache[home_id] = info
    return info


def member_can_add(pinfl, survey_uuid):
    """True = shu uyda ALLAQACHON bor (qo'shmaymiz)."""
    url = f"{API}/api/v1/survey_homes/family/member-can-add/check"
    code, txt = _req("POST", url, {"pinfl": int(pinfl), "survey_uuid": survey_uuid})
    if code != 200:
        return None  # noaniq — davom etamiz
    try:
        return bool(json.loads(txt).get("data", {}).get("status"))
    except Exception:
        return None


def gcp_person(pinfl, dob):
    url = (f"{API}/api/v1/gcp/pinfl?pinfl={pinfl}"
           f"&birth_date={dob:%Y-%m-%d}&form_name=survey_homes_family")
    for attempt in range(GCP_RETRY):
        code, txt = _req("GET", url)
        if code == 200:
            try:
                pd = json.loads(txt).get("data", {}).get("data")
            except Exception:
                pd = None
            return pd
        if code in (429, 0, 500, 502, 503):   # tezlik/tarmoq — kutib qayta urinamiz
            time.sleep(3 * (attempt + 1))
            continue
        return None
    return None


def build_payload(owner_pinfl, member_pinfl, dob, rel_id, pd, owner_phone):
    cur = str(pd.get("current_document") or "").strip()
    m = re.match(r"^(.*?)(\d+)$", cur)
    if not m:
        return None, "hujjat raqami yo'q"
    doc_serial = m.group(1).upper()
    doc_number = m.group(2)

    # hujjat turi: guvohnoma (type 80) -> 2, aks holda ID/pasport -> 1
    doc_type = "1"
    docs = pd.get("documents") or []
    matched = next((d for d in docs if str(d.get("document")) == cur), None)
    if matched is not None:
        doc_type = "2" if matched.get("type_id") == 80 else "1"
    elif docs and all(d.get("type_id") == 80 for d in docs):
        doc_type = "2"

    age = age_of(member_pinfl)
    # Telefon MAJBURIY EMAS. Uy egasining raqamini qayta-qayta ishlatib bo'lmaydi
    # (sayt "bu raqam N marta ishlatilgan" deb rad etadi), shuning uchun bo'sh qoldiramiz.
    phone = ""

    payload = {
        "pinfl": int(member_pinfl),
        "full_name": pd.get("full_name") or "",
        "address": None,
        "phone": phone,
        "birth_date": f"{dob:%Y-%m-%d}",
        "study_level_id": study_level(age, member_pinfl),
        "doc_serial": doc_serial,
        "doc_number": doc_number,
        "relationship": str(rel_id),
        "document_type": doc_type,
    }
    return payload, None


def save_member(owner_pinfl, survey_uuid, payload):
    q = "&".join(f"{k}={v}" for k, v in PARAMS.items())
    url = (f"{API}/web/v1/forms/survey_homes_family?_target=modal"
           f"&pinfl={owner_pinfl}&{q}&survey_uuid={survey_uuid}&owner_pinfl={owner_pinfl}")
    code, txt = _req("POST", url, payload)
    if code in (200, 201):
        # 200 bo'lsa ham ichida xato bo'lishi mumkin
        low = txt.lower()
        if code == 201 or ('"error":null' in low.replace(" ", "") or '"id"' in low):
            if "мавжуд" in txt or "mavjud" in txt or "biriktir" in txt.lower():
                return False, deep_message(txt)
            return True, None
        return False, deep_message(txt)
    return False, deep_message(txt)


# --- base.xlsx o'qish -----------------------------------------------------

def read_rows():
    wb = openpyxl.load_workbook(BASE_IN, data_only=True, read_only=True)
    ws = wb[SHEET]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        owner = ws.cell(r, 3).value      # C
        home_id = ws.cell(r, 5).value    # E
        member = ws.cell(r, 6).value     # F
        sana = ws.cell(r, 7).value       # G
        rel = ws.cell(r, 8).value        # H
        if owner is None and member is None:
            continue
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rows.append({
            "excel": vals,
            "owner": str(owner).strip() if owner is not None else "",
            "home_id": str(home_id).strip() if home_id is not None else "",
            "member": str(member).strip() if member is not None else "",
            "sana": sana,
            "rel": rel,
        })
    wb.close()
    return headers, rows


def key_of(row):
    return f"{row['home_id']}:{row['member']}"


# --- Natija Excel ---------------------------------------------------------

def write_excel(headers, rows, progress):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    out_headers = list(headers) + ["Natija"]
    ws.append(out_headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    styles = {
        "done": (PatternFill("solid", fgColor="D5F0DD"), Font(bold=True, color="1E7B3A")),
        "dup":  (PatternFill("solid", fgColor="DDE7F5"), Font(bold=True, color="1F4E78")),
        "error": (PatternFill("solid", fgColor="FDE9E7"), Font(bold=True, color="C0392B")),
        "pending": (PatternFill("solid", fgColor="FFF7E6"), Font(color="8A6D00")),
    }
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border

    for row in rows:
        st = progress.get(key_of(row), {})
        kind = st.get("kind", "pending")
        note = st.get("note", "Hali qilinmagan")
        ws.append(list(row["excel"]) + [note])
        r = ws.max_row
        for c in range(1, len(out_headers) + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = left if c in (2, len(out_headers)) else center
        fill, font = styles.get(kind, styles["pending"])
        last = ws.cell(r, len(out_headers))
        last.fill, last.font = fill, font

    widths = [6, 30, 16, 14, 10, 16, 13, 14, 46]
    for i in range(1, len(out_headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i <= len(widths) else 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(out_headers))}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    wb.save(OUT)


# --- Asosiy ---------------------------------------------------------------

def process_member(row):
    """Bitta a'zoni qo'shishga urinadi -> (kind, note). AuthError yuqoriga uzatiladi."""
    member = row["member"]
    dob = decode_birth(member)
    if not dob:
        return "error", "JSHSHIR noto'g'ri"

    rel_id = REL_MAP.get(norm_rel(row["rel"]))
    if rel_id is None:
        return "error", f"qarindoshlik aniqlanmadi: {row['rel']}"

    survey_uuid, owner_pinfl, owner_phone = get_home(row["home_id"])
    if not survey_uuid or not owner_pinfl:
        return "error", "uy survey_uuid topilmadi"

    if member_can_add(member, survey_uuid) is True:
        return "dup", "Allaqachon shu uyda bor"

    pd = gcp_person(member, dob)
    if not pd or pd.get("birth_date") == "-":
        return "error", "GCP: ma'lumot topilmadi (sana/JSHSHIR)"

    payload, err = build_payload(owner_pinfl, member, dob, rel_id, pd, owner_phone)
    if err:
        return "error", err

    ok, err = save_member(owner_pinfl, survey_uuid, payload)
    if ok:
        return "done", "Qo'shildi"
    return "error", err or "saqlash xatosi"


def main():
    # argumentlar:
    #   (bo'sh) -> 3 ta sinov | son -> shuncha | all -> hammasi
    #   member <JSHSHIR> -> faqat shu a'zo (sinov uchun)
    arg = sys.argv[1] if len(sys.argv) > 1 else "3"

    headers, rows = read_rows()
    progress = {}
    if os.path.exists(PROGRESS):
        progress = json.load(open(PROGRESS, encoding="utf-8"))

    if arg.lower() == "member":
        target = sys.argv[2]
        todo = [r for r in rows if r["member"] == target]
        print(f"Aniq a'zo rejimi: {target} -> {len(todo)} ta qator")
    else:
        limit = None if arg.lower() == "all" else int(arg)
        done_kinds = {"done", "dup", "error"}
        todo = [r for r in rows if progress.get(key_of(r), {}).get("kind") not in done_kinds]
        print(f"Jami {len(rows)} a'zo. Tayyor/skip: {len(rows)-len(todo)}. Qolgan: {len(todo)}.")
        if limit is not None:
            todo = todo[:limit]
            print(f"Bu safar {len(todo)} ta (sinov). Hammasi uchun: python add_family.py all")

    def set_status(row, kind, note):
        progress[key_of(row)] = {"kind": kind, "note": note}
        json.dump(progress, open(PROGRESS, "w", encoding="utf-8"), ensure_ascii=False, indent=2)

    counts = {"done": 0, "dup": 0, "error": 0}
    try:
        for i, row in enumerate(todo, 1):
            tag = f"[{i}/{len(todo)}] uy {row['home_id']} a'zo {row['member']} ({row['rel']})"
            kind = note = None
            for attempt in (1, 2):
                try:
                    kind, note = process_member(row)
                    break
                except AuthError:
                    if attempt == 2:
                        print("\n!!! Token yangilanmadi (401). To'xtatildi. "
                              "Keyinroq qayta ishga tushiring — qolganidan davom etadi.")
                        return
                    refresh_token()
                except Exception as e:
                    kind, note = "error", f"kutilmagan xato: {str(e)[:150]}"
                    break
            set_status(row, kind, note)
            counts[kind] = counts.get(kind, 0) + 1
            label = {"done": "QO'SHILDI", "dup": "ALLAQACHON bor", "error": "XATO"}[kind]
            print(f"{tag} -> {label}" + (f": {note[:80]}" if kind == "error" else ""))
            if i % EXCEL_EVERY == 0:
                write_excel(headers, rows, progress)
            time.sleep(DELAY)
    finally:
        write_excel(headers, rows, progress)
        print(f"\nYakun: qo'shildi {counts['done']}, allaqachon {counts['dup']}, "
              f"xato {counts['error']}.  Natija -> {OUT}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
    main()
