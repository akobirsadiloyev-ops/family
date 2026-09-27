"""
family_search.xlsx -> C ustunidagi JSHSHIR asosida E ustunga "Tug'ilgan sana" qo'shadi.

O'zbekiston JSHSHIR (14 raqam) tarkibida tug'ilgan sana kodlangan:
  1-raqam  -> asr + jins (1,2=18xx; 3,4=19xx; 5,6=20xx)
  2-7 raqam -> KK OO YY (kun, oy, yil)
Sayt ham aynan shu qiymatni ko'rsatadi, shuning uchun natija bir xil.

Ishlatish:
    python add_birthdate.py
"""

import copy
import datetime

import openpyxl

FILE = "family_search.xlsx"
JSHSHIR_COL = "C"   # JSHSHIR ustuni
NEW_COL = "E"       # yangi "Tug'ilgan sana" ustuni
NEW_HEADER = "Tug'ilgan sana"


def decode_birthdate(pinfl):
    if pinfl is None:
        return None
    p = str(pinfl).strip()
    if len(p) != 14 or not p.isdigit():
        return None
    century = int(p[0])
    dd, mm, yy = int(p[1:3]), int(p[3:5]), int(p[5:7])
    base = {1: 1800, 2: 1800, 3: 1900, 4: 1900, 5: 2000, 6: 2000}.get(century)
    if base is None:
        return None
    try:
        return datetime.date(base + yy, mm, dd)
    except ValueError:
        return None


def main():
    wb = openpyxl.load_workbook(FILE)
    ws = wb.active

    # --- Sarlavha (mavjud sarlavha stilini nusxalab) ---
    src_hdr = ws["D1"]
    hc = ws[f"{NEW_COL}1"]
    hc.value = NEW_HEADER
    hc.font = copy.copy(src_hdr.font)
    hc.fill = copy.copy(src_hdr.fill)
    hc.alignment = copy.copy(src_hdr.alignment)
    hc.border = copy.copy(src_hdr.border)

    # --- Qiymatlar ---
    filled, missing = 0, 0
    for row in range(2, ws.max_row + 1):
        jshshir = ws[f"{JSHSHIR_COL}{row}"].value
        d = decode_birthdate(jshshir)
        cell = ws[f"{NEW_COL}{row}"]
        # ma'lumot katakcha stilini nusxalash (chegara/tekislash)
        src = ws[f"D{row}"]
        cell.border = copy.copy(src.border)
        cell.alignment = copy.copy(src.alignment)
        cell.number_format = "@"  # matn -> sana ko'rinishi buzilmasin
        if d:
            cell.value = d.strftime("%d.%m.%Y")
            filled += 1
        else:
            cell.value = ""
            missing += 1

    # --- Ustun kengligi, filtr ---
    ws.column_dimensions[NEW_COL].width = 15
    ws.auto_filter.ref = f"A1:{NEW_COL}{ws.max_row}"

    wb.save(FILE)
    print(f"Tayyor: {filled} ta sana yozildi, {missing} ta bo'sh -> {FILE}")


if __name__ == "__main__":
    main()
