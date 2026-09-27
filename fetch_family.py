"""
family_search.xlsx -> ins.ihma.uz dan har bir odamning oila a'zolarini olib,
family_members.xlsx ga yozadi.

  C ustun -> JSHSHIR,  E ustun -> Tug'ilgan sana (DD.MM.YYYY)

Qanday ishlaydi:
  1. Brauzer ochiladi (sessiya .browser_profile papkasida saqlanadi).
  2. Login qilasiz va ochilgan sahifada BITTA odamni qo'lda qidirasiz
     ("Qidiruv" tugmasi). Skript shu so'rovni ushlab, hujjat turini oladi.
  3. Qolgan hammasi avtomatik: saytning o'z API'si
     (/api/Person/GetPersonFamily) orqali so'raladi.

Natija har qatordan keyin family_results.json ga saqlanadi — skript to'xtasa,
qayta ishga tushiring, tayyor bo'lganlar o'tkazib yuboriladi.

Ishlatish:
    pip install playwright openpyxl
    python -m playwright install chromium
    python fetch_family.py
"""

import datetime
import json
import os
import sys
import time

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

from add_birthdate import decode_birthdate

# --- Sozlamalar ----------------------------------------------------------
IN_FILE = "family_search.xlsx"
OUT_FILE = "family_members.xlsx"
RESULTS_FILE = "family_results.json"
PROFILE_DIR = ".browser_profile"

NUM_COL, NAME_COL, JSHSHIR_COL, DOB_COL = "A", "B", "C", "E"

PAGE_URL = (
    "https://ins.ihma.uz/documents/ChPnsApplication/edit/0"
    "?fromRegistration=13628469&appId=26"
    "&appName=%D0%91%D0%BE%D0%BB%D0%B0%D0%BB%D0%B0%D1%80+%D0%BD%D0%B0%D1%84"
    "%D0%B0%D2%9B%D0%B0%D1%81%D0%B8+%D1%91%D0%BA%D0%B8+%D0%BC%D0%BE%D0%B4%D0"
    "%B4%D0%B8%D0%B9+%D1%91%D1%80%D0%B4%D0%B0%D0%BC+%D1%82%D0%B0%D0%B9%D0%B8"
    "%D0%BD%D0%BB%D0%B0%D1%88&appCode=4001&appCategoryId=4"
)
API = "https://ins.ihma.uz/api"
DELAY = 0.4  # so'rovlar orasidagi pauza (soniya)
# -------------------------------------------------------------------------

# Brauzer ichida ishlaydi: OIDC tokenni sessionStorage dan olib, API ga POST qiladi
JS_POST = """
async ({url, body, headers}) => {
  let token = null;
  for (let i = 0; i < sessionStorage.length; i++) {
    const k = sessionStorage.key(i);
    if (k.startsWith('oidc.user:')) {
      try { token = JSON.parse(sessionStorage.getItem(k)).access_token; } catch (e) {}
    }
  }
  const h = {...headers, 'Content-Type': 'application/json'};
  if (token) h['Authorization'] = 'Bearer ' + token;
  const r = await fetch(url, {method: 'POST', headers: h, body: JSON.stringify(body)});
  return {status: r.status, text: await r.text()};
}
"""

# ushlangan so'rovdan qayta ishlatiladigan sarlavhalar
KEEP_HEADERS = {"accept", "accept-language", "x-requested-with"}


def read_people():
    ws = openpyxl.load_workbook(IN_FILE, read_only=True).active
    people = []
    for row in range(2, ws.max_row + 1):
        pinfl = ws[f"{JSHSHIR_COL}{row}"].value
        if pinfl is None or not str(pinfl).strip():
            continue
        pinfl = str(pinfl).strip()
        dob = ws[f"{DOB_COL}{row}"].value
        if isinstance(dob, (datetime.date, datetime.datetime)):
            dob = dob.strftime("%d.%m.%Y")
        elif not dob:
            d = decode_birthdate(pinfl)
            dob = d.strftime("%d.%m.%Y") if d else ""
        people.append({
            "num": ws[f"{NUM_COL}{row}"].value,
            "name": ws[f"{NAME_COL}{row}"].value or "",
            "pinfl": pinfl,
            "dob": str(dob).strip(),
        })
    return people


def load_results():
    if os.path.exists(RESULTS_FILE):
        with open(RESULTS_FILE, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_results(results):
    tmp = RESULTS_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    os.replace(tmp, RESULTS_FILE)


def error_text(resp):
    try:
        data = json.loads(resp["text"])
        if isinstance(data, dict):
            for key in ("message", "Message", "error", "title", "detail"):
                if data.get(key):
                    return f"HTTP {resp['status']}: {data[key]}"
    except (ValueError, TypeError):
        pass
    return f"HTTP {resp['status']}: {resp['text'][:200]}"


def fmt_date(value):
    """'1991-11-30T00:00:00' yoki '30.11.1991' -> '30.11.1991'."""
    if not value:
        return ""
    s = str(value)
    try:
        return datetime.date.fromisoformat(s[:10]).strftime("%d.%m.%Y")
    except ValueError:
        return s


def fmt_status(value):
    if isinstance(value, dict):
        return value.get("name") or value.get("text") or ""
    return value or ""


def wait_for_template(page):
    """Foydalanuvchi qo'lda qidiruv qilguncha kutadi, so'rov tanasini qaytaradi."""
    captured = {}

    def on_request(req):
        if req.method == "POST" and "/api/Person/GetByPassportData" in req.url:
            try:
                captured["body"] = req.post_data_json
                captured["headers"] = {
                    k: v for k, v in req.headers.items() if k.lower() in KEEP_HEADERS
                }
            except Exception:
                pass

    page.on("request", on_request)
    print("\n>>> Brauzerda login qiling, so'ng ochilgan sahifada BITTA odamni")
    print(">>> qo'lda qidiring (Hujjat turi + Tug'ilgan sana + JSHSHIR -> Qidiruv).")
    print(">>> Skript shu so'rovni ushlab olgach, avtomatik davom etadi...\n")
    while "body" not in captured:
        page.wait_for_timeout(500)
    page.remove_listener("request", on_request)
    return captured["body"], captured["headers"]


def post(page, path, body, headers):
    while True:
        resp = page.evaluate(JS_POST, {"url": f"{API}{path}", "body": body, "headers": headers})
        if resp["status"] != 401:
            return resp
        input("\n!!! Sessiya tugagan (401). Brauzerda qayta login qiling va Enter bosing...")


def fetch_family(page, template, headers, person):
    body = {**template, "pinfl": person["pinfl"], "dateOfBirth": person["dob"],
            "includePhoto": False}

    resp = post(page, "/Person/GetPersonFamily", body, headers)
    if resp["status"] != 200:
        # saytdagidek: avval shaxsni topib, keyin oilani so'raymiz
        first = post(page, "/Person/GetByPassportData", body, headers)
        if first["status"] != 200:
            return {"status": "error", "error": error_text(first)}
        resp = post(page, "/Person/GetPersonFamily", body, headers)
        if resp["status"] != 200:
            return {"status": "error", "error": error_text(resp)}

    data = json.loads(resp["text"])
    family = data.get("personFamily") if isinstance(data, dict) else None
    members = [
        {
            "fullName": m.get("fullName") or "",
            "pinfl": m.get("pinfl") or "",
            "birthOn": fmt_date(m.get("birthOn")),
            "status": fmt_status(m.get("familyMemberStatus")),
        }
        for m in (family or [])
    ]
    return {"status": "ok", "members": members}


def write_excel(people, results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Oila a'zolari"

    headers = ["№", "Uy egasi (F.I.O.)", "Uy egasi JSHSHIR", "Oila a'zosi (F.I.O.)",
               "JSHSHIR", "Tug'ilgan sana", "Holati", "Izoh"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    group_fill = PatternFill("solid", fgColor="EEF3F8")  # har ikkinchi oila foni
    err_font = Font(color="C0392B")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border

    ok = err = pending = n_members = 0
    for g, p in enumerate(people):
        res = results.get(p["pinfl"])
        if res is None:
            lines, note = [["", "", "", "", ""]], "Hali so'ralmagan"
            pending += 1
        elif res["status"] == "ok":
            lines = [[m["fullName"], m["pinfl"], m["birthOn"], m["status"], ""]
                     for m in res["members"]] or [["", "", "", "", "Oila a'zolari topilmadi"]]
            note = None
            ok += 1
            n_members += len(res["members"])
        else:
            lines, note = [["", "", "", "", res["error"]]], None
            err += 1

        for line in lines:
            if note:
                line[-1] = note
            ws.append([p["num"], p["name"], p["pinfl"], *line])
            r = ws.max_row
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=r, column=col)
                c.border = border
                c.alignment = left if col in (2, 4, 8) else center
                if g % 2:
                    c.fill = group_fill
            for col in (3, 5, 6):
                ws.cell(row=r, column=col).number_format = "@"
            if line[-1]:
                ws.cell(row=r, column=8).font = err_font

    widths = [6, 34, 17, 34, 17, 14, 16, 40]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False

    wb.save(OUT_FILE)
    print(f"\nExcel: {OUT_FILE}  (muvaffaqiyatli: {ok}, xato: {err}, "
          f"so'ralmagan: {pending}, jami oila a'zosi: {n_members})")


def main():
    people = read_people()
    results = load_results()
    todo = [p for p in people if results.get(p["pinfl"], {}).get("status") != "ok"]
    print(f"Excel'da {len(people)} ta odam. Tayyor: {len(people) - len(todo)}, "
          f"qolgan: {len(todo)}")

    if todo:
        with sync_playwright() as pw:
            ctx = pw.chromium.launch_persistent_context(
                PROFILE_DIR, headless=False, viewport=None, args=["--start-maximized"])
            page = ctx.pages[0] if ctx.pages else ctx.new_page()
            page.goto(PAGE_URL)
            template, headers = wait_for_template(page)
            print(f"So'rov ushlandi: {json.dumps(template, ensure_ascii=False)}")

            try:
                for i, p in enumerate(todo, start=1):
                    if not p["dob"]:
                        results[p["pinfl"]] = {"status": "error",
                                               "error": "Tug'ilgan sana yo'q"}
                    else:
                        try:
                            results[p["pinfl"]] = fetch_family(page, template, headers, p)
                        except Exception as e:  # tarmoq/brauzer xatosi
                            results[p["pinfl"]] = {"status": "error", "error": str(e)[:200]}
                    save_results(results)
                    res = results[p["pinfl"]]
                    info = (f"{len(res['members'])} ta a'zo" if res["status"] == "ok"
                            else f"XATO: {res['error']}")
                    print(f"[{i}/{len(todo)}] {p['pinfl']} {p['name']} -> {info}")
                    time.sleep(DELAY)
            except KeyboardInterrupt:
                print("\nTo'xtatildi. Qayta ishga tushirsangiz, qolganidan davom etadi.")
            finally:
                ctx.close()

    write_excel(people, results)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
